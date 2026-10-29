"""Staged operational XLSX import. Preview is immutable input for atomic commit."""
import hashlib
import io
import json
import math
import re
from sqlalchemy.exc import IntegrityError
import zipfile
from collections import Counter
from contextlib import contextmanager
from datetime import datetime, timedelta
from pathlib import PurePosixPath
from uuid import uuid4

from defusedxml.ElementTree import iterparse
from fastapi import Depends, HTTPException, Query, Request
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import Field, StrictInt, ValidationError
from starlette.concurrency import run_in_threadpool

from .domain import MODEL_VERSION, score_case, target_team, timestamp
from .scoring import capture_origin
from .policy import current_policy,rule_version
from .schemas import CaseCreate, Input, Kind
from .service import audit

MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
MAX_BYTES, MAX_ROWS, MAX_COLS = 5 * 1024 * 1024, 2000, 64
FIELDS = {
    'source_ref': ('ID referensi sumber', ['idsumber', 'idkasus', 'idlead'], 'Teks unik per alur, maks. 80 karakter. Simpan sebagai teks agar nol awal tidak hilang.'),
    'customer': ('Alias demo', ['namademo', 'aliasdemo', 'iddebitur'], 'Alias sintetis 2–100 karakter.'),
    'region_id': ('Wilayah', ['kodewilayah', 'wilayah', 'kotadomisili'], 'Kode atau nama wilayah yang terdaftar.'),
    'difficulty': ('Tingkat kesulitan', ['tingkatkesulitan', 'kategoricase'], 'Low, Medium, atau High.'),
    'amount': ('Nominal rupiah', ['nominaltagihan', 'tagihanpenuh', 'nominaldiajukan', 'outstanding'], 'Rupiah bulat > 0 sampai 1 miliar. Tanpa Rp, titik, atau koma pada teks.'),
    'dpd': ('Hari keterlambatan', ['dpdhari'], 'Bilangan bulat 0–3650. Kosong/tidak dipetakan = 0.'),
    'income': ('Penghasilan bulanan', ['penghasilanbulanan'], 'Rupiah bulat > 0 sampai 1 miliar.'),
    'interest': ('Minat', ['minat'], 'Bilangan bulat 1–5.'),
    'documents_complete': ('Dokumen lengkap', ['dokumenlengkap'], 'Ya/Tidak, true/false, atau 1/0. Kosong = Tidak.'),
}


def fields(kind):
    names = ['source_ref', 'customer', 'region_id', 'difficulty', 'amount'] + (['dpd'] if kind == 'collection' else ['income', 'interest', 'documents_complete'])
    return [dict(name=n, label=FIELDS[n][0], required=n not in {'dpd', 'documents_complete'}, help=FIELDS[n][2]) for n in names]


class PreviewInput(Input):
    sheet: str = Field(min_length=1, max_length=31)
    header_row: int = Field(default=1, ge=1, le=20, strict=True)
    mapping: dict[str, StrictInt] = Field(max_length=9)


class CommitInput(Input):
    expected_revision: int = Field(ge=1, strict=True)
    acknowledge_ignored: bool = False


def normalized(value):
    return re.sub(r'[^a-z0-9]', '', str(value or '').lower())


def inspect_package(content):
    if not content or len(content) > MAX_BYTES:
        raise HTTPException(413, 'Berkas harus berisi data dan maksimal 5 MB.')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            names = {e.filename for e in entries}
            if not {'[Content_Types].xml', 'xl/workbook.xml'} <= names:
                raise ValueError()
            if len(entries) > 2000 or sum(e.file_size for e in entries) > 30 * 1024 * 1024:
                raise HTTPException(413, 'Isi XLSX terkompresi terlalu besar.')
            if any(e.flag_bits & 1 or 'vbaproject' in e.filename.lower() or e.filename.startswith(('xl/externalLinks/', 'xl/embeddings/')) for e in entries):
                raise HTTPException(422, 'Makro, berkas terenkripsi, atau tautan workbook eksternal tidak didukung.')
            count = 0
            for entry in entries:
                if not re.fullmatch(r'xl/worksheets/sheet\d+\.xml', entry.filename):
                    continue
                with archive.open(entry) as stream:
                    for _, cell in iterparse(stream, events=['end']):
                        if cell.tag.endswith('}c'):
                            if any(child.tag.rsplit('}', 1)[-1] in {'v', 'is', 'f'} for child in cell):
                                match = re.fullmatch(r'([A-Z]+)(\d+)', cell.attrib.get('r', ''))
                                if not match or column_index_from_string(match[1]) > MAX_COLS or int(match[2]) > MAX_ROWS + 20:
                                    raise HTTPException(413, 'Batas workbook: 64 kolom dan 2.000 baris data, header pada baris 1–20.')
                                count += 1
                                if count > 200000:
                                    raise HTTPException(413, 'Workbook melebihi 200.000 sel berisi.')
                            cell.clear()
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, 'Berkas bukan XLSX yang valid atau tidak dapat dibaca.')


@contextmanager
def workbook(content):
    book = None
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if not 1 <= len(book.worksheets) <= 20:
            raise HTTPException(413, 'Workbook harus memiliki 1–20 worksheet.')
        yield book
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, 'Workbook tidak dapat dibaca. Gunakan XLSX tanpa formula atau makro.')
    finally:
        if book:
            book.close()


def sheet_layout(book, kind, sheet, header_row):
    if sheet not in book.sheetnames:
        raise HTTPException(422, 'Worksheet tidak ditemukan.')
    ws = book[sheet]
    ws.reset_dimensions()
    row = next(ws.iter_rows(min_row=header_row, max_row=header_row, max_col=MAX_COLS))
    if any(c.data_type == 'f' for c in row):
        raise HTTPException(422, 'Header tidak boleh berisi formula.')
    last = max((i + 1 for i, c in enumerate(row) if c.value is not None), default=0)
    if not last:
        raise HTTPException(422, 'Baris header kosong. Pilih baris header yang benar.')
    for values in ws.iter_rows(min_row=header_row + 1, max_row=MAX_ROWS + 20, max_col=MAX_COLS, values_only=True):
        last = max(last, max((i + 1 for i, value in enumerate(values) if value is not None), default=0))
    columns = [dict(index=i + 1, letter=get_column_letter(i + 1), header=str(c.value or '')[:120]) for i, c in enumerate(row[:last])]
    mapping = {}
    for f in fields(kind):
        aliases = {normalized(f['name']), *FIELDS[f['name']][1]}
        matches = [c['index'] for c in columns if normalized(c['header']) in aliases]
        if len(matches) == 1:
            mapping[f['name']] = matches[0]
    return dict(sheet=sheet, header_row=header_row, columns=columns, mapping=mapping, fields=fields(kind))


def metadata(content, kind):
    with workbook(content) as book:
        sheets = [dict(name=s.title, hidden=s.sheet_state != 'visible') for s in book.worksheets]
        best, score = None, -1
        for sheet in book.worksheets:
            for header_row, row in enumerate(sheet.iter_rows(max_row=20, max_col=MAX_COLS), 1):
                values = {normalized(c.value) for c in row if c.data_type != 'f'}
                matched = sum(bool(values & {normalized(f['name']), *FIELDS[f['name']][1]}) for f in fields(kind))
                if matched > score:
                    best, score = (sheet.title, header_row), matched
        legacy = any('kamus' in s['name'].lower() or '1200' in s['name'] for s in sheets)
        return dict(sheets=sheets, suggested_sheet=best[0], suggested_header_row=best[1],
                    warnings=['Template pengumpulan/historis terdeteksi. Impor ini hanya membuat kasus baru; histori, pembayaran, dan target tidak diimpor.'] if legacy else [],
                    max_rows=MAX_ROWS, max_bytes=MAX_BYTES)


def require_batch(conn, batch_id, user, now, allow_expired=False):
    row = conn.execute('SELECT * FROM import_batches WHERE id=? AND actor_id=?', (batch_id, user['id'])).fetchone()
    if row is None:
        raise HTTPException(404, 'Batch impor tidak ditemukan.')
    batch = dict(row)
    if (batch['state'] == 'discarded' or (batch['state'] != 'committed' and batch['expires_at'] <= timestamp(now))) and not allow_expired:
        raise HTTPException(410, 'Draft impor kedaluwarsa. Unggah ulang workbook.')
    return batch


def batch_result(batch, now, offset=0, limit=25):
    result = {k: batch[k] for k in ['id', 'kind', 'filename', 'file_sha256', 'created_at', 'expires_at', 'state', 'revision']}
    result['metadata'] = json.loads(batch['metadata'])
    if batch['state'] not in {'committed', 'discarded'} and batch['expires_at'] <= timestamp(now):
        result['state'] = 'expired'
    if batch['preview'] and result['state'] != 'expired':
        p = json.loads(batch['preview'])
        rows = p.pop('rows')
        result['preview'] = dict(p, items=rows[offset:offset + limit], offset=offset, limit=limit)
    if batch['result']:
        committed = json.loads(batch['result'])
        items = committed.pop('items')
        result['result'] = dict(committed, items=items[offset:offset + limit], offset=offset, limit=limit)
    return result


def upload(db, content, kind, filename, user, now):
    filename = PurePosixPath(filename.replace('\\', '/')).name
    if not filename.lower().endswith('.xlsx'):
        raise HTTPException(422, 'Hanya berkas .xlsx yang didukung.')
    inspect_package(content)
    meta = metadata(content, kind)
    bid = 'IMP-' + uuid4().hex[:16].upper()
    with db.connect(write=True) as conn:
        conn.execute("UPDATE import_batches SET content=NULL WHERE state!='committed' AND expires_at<=?", (timestamp(now),))
        count = conn.execute("SELECT COUNT(*) FROM import_batches WHERE actor_id=? AND state IN ('uploaded','previewed') AND expires_at>?", (user['id'], timestamp(now))).fetchone()[0]
        if count >= 10:
            raise HTTPException(409, 'Maksimal 10 draft aktif. Selesaikan draft atau tunggu masa berlaku 24 jam.')
        order = conn.execute('SELECT COALESCE(MAX(upload_order),0)+1 FROM import_batches').fetchone()[0]
        conn.execute('INSERT INTO import_batches(id,actor_id,kind,filename,file_sha256,content,created_at,expires_at,metadata,upload_order) VALUES (?,?,?,?,?,?,?,?,?,?)',
                     (bid, user['id'], kind, filename, hashlib.sha256(content).hexdigest(), content, timestamp(now), timestamp(now + timedelta(hours=24)), json.dumps(meta), order))
        audit(conn, user, 'import_upload', None, dict(batch_id=bid, kind=kind, filename=filename, bytes=len(content)), now)
        return batch_result(require_batch(conn, bid, user, now), now)


def number(value):
    if isinstance(value, bool) or not isinstance(value, (str, int, float)):
        raise ValueError('Gunakan bilangan bulat tanpa Rp atau pemisah ribuan.')
    if isinstance(value, str):
        if not re.fullmatch(r'[0-9]+', value.strip()):
            raise ValueError('Gunakan bilangan bulat tanpa Rp atau pemisah ribuan.')
        return int(value.strip())
    if isinstance(value, float) and (not math.isfinite(value) or not value.is_integer()):
        raise ValueError('Nilai harus bilangan bulat.')
    return int(value)


def make_preview(content, kind, payload, conn):
    policy = current_policy(conn)
    with workbook(content) as book:
        layout = sheet_layout(book, kind, payload.sheet, payload.header_row)
        allowed = {f['name'] for f in fields(kind)}
        required = {f['name'] for f in fields(kind) if f['required']}
        mapping = payload.mapping
        indices = list(mapping.values())
        if set(mapping) - allowed or required - set(mapping):
            raise HTTPException(422, 'Petakan semua kolom wajib; hanya field operasional alur ini yang boleh dipetakan.')
        if len(set(indices)) != len(indices) or any(type(v) is not int or v < 1 or v > len(layout['columns']) for v in indices):
            raise HTTPException(422, 'Pemetaan kolom berulang atau indeks kolom tidak valid.')
        regions = {str(r['id']).casefold(): r['id'] for r in conn.execute('SELECT * FROM regions WHERE active=1')}
        regions.update({str(r['name']).casefold(): r['id'] for r in conn.execute('SELECT * FROM regions WHERE active=1')})
        existing = {r['source_key'] for r in conn.execute('SELECT source_key FROM import_case_refs WHERE kind=?', (kind,))}
        rows, skipped, pending_blank = [], 0, 0
        ws = book[payload.sheet]
        for row_number, cells in enumerate(ws.iter_rows(min_row=payload.header_row + 1, max_row=MAX_ROWS + 20, max_col=MAX_COLS), payload.header_row + 1):
            if all(c.value is None or str(c.value).strip() == '' for c in cells):
                pending_blank += 1; continue
            skipped += pending_blank; pending_blank = 0
            if len(rows) >= MAX_ROWS:
                raise HTTPException(413, 'Maksimal 2.000 baris data per batch; tidak ada baris yang diimpor.')
            errors, raw, data = [], {}, dict(kind=kind)
            for c in cells:
                if c.data_type in {'f', 'e'}:
                    errors.append(f'{c.coordinate}: formula atau error Excel tidak didukung.')
            for field, index in mapping.items():
                value = cells[index - 1].value
                raw[field] = str(value)[:200] if value is not None else ''
                if value is None or str(value).strip() == '':
                    if field in required:
                        errors.append(f'{field}: nilai wajib diisi.')
                    continue
                try:
                    if field in {'source_ref', 'customer', 'region_id', 'difficulty'}:
                        if not isinstance(value, str):
                            raise ValueError('Simpan sebagai teks, terutama ID referensi agar nol awal terjaga.')
                        value = value.strip()
                        if field == 'source_ref' and not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_.:/-]{0,79}', value):
                            raise ValueError('ID memakai 1–80 karakter huruf/angka dan _ . : / -.')
                        if field == 'region_id':
                            if value.casefold() not in regions:
                                raise ValueError('Wilayah tidak terdaftar.')
                            value = regions[value.casefold()]
                        if field == 'difficulty':
                            value = {'low': 'Low', 'medium': 'Medium', 'high': 'High'}.get(value.casefold(), value)
                    elif field == 'documents_complete':
                        booleans = {'ya': True, 'true': True, '1': True, 'tidak': False, 'false': False, '0': False}
                        if str(value).strip().casefold() not in booleans:
                            raise ValueError('Gunakan Ya/Tidak, true/false, atau 1/0.')
                        value = booleans[str(value).strip().casefold()]
                    else:
                        value = number(value)
                    data[field] = value
                except ValueError as exc:
                    errors.append(f'{field}: {exc}')
            ref = data.pop('source_ref', '')
            if ref.upper() in existing:
                errors.append('source_ref: sudah pernah diimpor pada alur ini.')
            try:
                case = CaseCreate.model_validate(data).model_dump()
                score = score_case(case,policy=policy['config'])[0]
                case.update(score=score)
                team = target_team({**case,'medium_threshold':policy['config']['medium_threshold']})
            except ValidationError as exc:
                errors.extend(f"{'.'.join(str(x) for x in e['loc']) or 'baris'}: {e['msg']}" for e in exc.errors())
                case, team = None, None
            rows.append(dict(row=row_number, source_ref=ref, raw=raw, data=case, team=team, errors=list(dict.fromkeys(errors))))
        counts = Counter(r['source_ref'].upper() for r in rows if r['source_ref'])
        for r in rows:
            if counts[r['source_ref'].upper()] > 1:
                r['errors'].append('source_ref: ID berulang di workbook.')
        invalid = sum(bool(r['errors']) for r in rows)
        ignored = [c for c in layout['columns'] if c['index'] not in indices]
        return dict(sheet=payload.sheet, header_row=payload.header_row, mapping=mapping, columns=layout['columns'],
                    ignored_columns=ignored, score_version=rule_version(policy['revision']), policy_revision=policy['revision'], rows=rows,
                    total=len(rows), valid=len(rows) - invalid, invalid=invalid,
                    can_commit=bool(rows) and not invalid, skipped_blank=skipped,
                    amount=sum(r['data']['amount'] for r in rows if not r['errors'] and r['data']))


def preview(db, bid, payload, user, now):
    with db.connect() as conn:
        batch = require_batch(conn, bid, user, now)
        if batch['state'] == 'committed':
            raise HTTPException(409, 'Batch sudah diimpor; gunakan workbook baru untuk perubahan.')
        result = make_preview(batch['content'], batch['kind'], payload, conn)
    with db.connect(write=True) as conn:
        latest = require_batch(conn, bid, user, now)
        if latest['revision'] != batch['revision'] or latest['state'] == 'committed':
            raise HTTPException(409, 'Draft berubah. Muat ulang sebelum preview ulang.')
        conn.execute("UPDATE import_batches SET state='previewed',revision=revision+1,preview=? WHERE id=?", (json.dumps(result), bid))
        audit(conn, user, 'import_preview', None, dict(batch_id=bid, total=result['total'], invalid=result['invalid'], sheet=payload.sheet), now)
        return batch_result(require_batch(conn, bid, user, now), now)


def commit(db, bid, payload, user, now):
    try:
        with db.connect(write=True) as conn:
            batch = require_batch(conn, bid, user, now)
            if batch['revision'] != payload.expected_revision:
                raise HTTPException(409, 'Preview berubah. Muat ulang dan tinjau preview terbaru.')
            if batch['state'] == 'committed':
                return dict(batch_result(batch, now), replayed=True)
            p = json.loads(batch['preview'] or '{}')
            policy = current_policy(conn)
            if not p.get('can_commit') or p.get('score_version') != rule_version(policy['revision']):
                raise HTTPException(409, 'Preview belum valid atau aturan berubah. Jalankan preview ulang.')
            if p['ignored_columns'] and not payload.acknowledge_ignored:
                raise HTTPException(422, 'Konfirmasi kolom yang tidak dipetakan sebelum impor.')
            items = []
            policy = current_policy(conn)
            if p.get('policy_revision',0)!=policy['revision']:
                raise HTTPException(409,'Kebijakan berubah. Jalankan preview impor ulang.')
            for row in p['rows']:
                case = CaseCreate.model_validate({k: v for k, v in row['data'].items() if k != 'score'}).model_dump()
                if not conn.execute('SELECT 1 FROM regions WHERE id=? AND active=1', (case['region_id'],)).fetchone():
                    raise HTTPException(409, 'Master wilayah berubah. Jalankan preview ulang.')
                if conn.execute('SELECT 1 FROM import_case_refs WHERE kind=? AND source_key=?', (batch['kind'], row['source_ref'].upper())).fetchone():
                    raise HTTPException(409, f"Referensi {row['source_ref']} sudah diimpor batch lain. Tidak ada kasus yang disimpan; preview ulang.")
                case['documents_complete'] = int(case['documents_complete'])
                case.update(id=f'{"TAG" if batch["kind"] == "collection" else "AJU"}-{uuid4().hex[:10].upper()}',
                            created_at=timestamp(now), score=score_case(case,policy=policy['config'])[0], score_version=rule_version(policy['revision']), source='excel-operational')
                conn.execute(f'INSERT INTO cases ({",".join(case)}) VALUES ({",".join("?" for _ in case)})', tuple(case.values()))
                capture_origin(conn,case)
                conn.execute('INSERT INTO import_case_refs VALUES (?,?,?,?,?)', (batch['kind'], row['source_ref'].upper(), row['source_ref'], case['id'], bid))
                audit(conn, user, 'import_case', case['id'], dict(batch_id=bid, source_ref=row['source_ref']), now)
                items.append(dict(case_id=case['id'], source_ref=row['source_ref'], customer=case['customer']))
            result = dict(total=len(items), committed_at=timestamp(now), items=items)
            conn.execute("UPDATE import_batches SET state='committed',result=?,content=NULL WHERE id=?", (json.dumps(result), bid))
            audit(conn, user, 'import_commit', None, dict(batch_id=bid, rows=len(items), kind=batch['kind'], file_sha256=batch['file_sha256'], sheet=p['sheet'], mapping=p['mapping']), now)
            return dict(batch_result(require_batch(conn, bid, user, now), now), replayed=False)
    except IntegrityError:
        raise HTTPException(409, 'Konflik data impor. Seluruh batch dibatalkan; jalankan preview ulang.')
    except ValidationError:
        raise HTTPException(409, 'Validasi kasus berubah. Seluruh batch dibatalkan; jalankan preview ulang.')


def template(kind, regions):
    book = Workbook()
    ws = book.active; ws.title = 'Data'
    specs = fields(kind)
    ws.append([f['name'] for f in specs])
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f'A1:{get_column_letter(len(specs))}2001'
    ws.sheet_view.showGridLines = False
    for i, f in enumerate(specs, 1):
        cell = ws.cell(1, i); cell.font = Font(color='FFFFFF', bold=True); cell.fill = PatternFill('solid', fgColor='087E8B')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(i)].width = max(20, len(f['name']) + 4)
    ws.row_dimensions[1].height = 30
    lists = {'difficulty': 'Low,Medium,High', 'region_id': ','.join(r['id'] for r in regions), 'documents_complete': 'Ya,Tidak'}
    for i, f in enumerate(specs, 1):
        if f['name'] in lists:
            validation = DataValidation(type='list', formula1='"' + lists[f['name']] + '"', allow_blank=not f['required'])
            validation.errorTitle = 'Nilai tidak sesuai'; validation.error = f['help']; validation.showErrorMessage = True
            ws.add_data_validation(validation); validation.add(f'{get_column_letter(i)}2:{get_column_letter(i)}2001')
        if f['name'] == 'source_ref':
            ws.column_dimensions[get_column_letter(i)].number_format = '@'
    guide = book.create_sheet('Petunjuk')
    guide.append(['Template operasional ' + ('penagihan' if kind == 'collection' else 'pengajuan')])
    guide.append(['Isi worksheet Data mulai baris 2. Data belum berisi contoh yang dapat terimpor.'])
    guide.append(['Satu baris = satu kasus baru. Gunakan identitas sintetis. ID sumber unik per alur; impor tidak memperbarui kasus lama.'])
    guide.append(['Nominal menjadi tagihan penuh/nominal pengajuan. Tidak mengimpor pembayaran, petugas, histori, atau target.'])
    guide.append(['Batas 5 MB, 2.000 baris, 64 kolom. Tanpa formula/makro. Preview wajib sebelum commit.'])
    guide.append(['Wilayah tersedia: ' + ', '.join(f'{r["id"]} ({r["name"]})' for r in regions)])
    guide.append(['Kolom', 'Wajib', 'Aturan', 'Contoh sintetis'])
    examples = dict(source_ref='DEMO-0001', customer='Alias Sintetis 001', region_id=regions[0]['id'] if regions else 'SMG',
                    difficulty='Medium', amount=1000000, dpd=12, income=5000000, interest=4, documents_complete='Tidak')
    for f in specs:
        guide.append([f['name'], 'Ya' if f['required'] else 'Tidak', f['help'], examples[f['name']]])
    for col, width in [('A', 48), ('B', 12), ('C', 75), ('D', 25)]:
        guide.column_dimensions[col].width = width
    for row in guide:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        guide.row_dimensions[row[0].row].height = 44
    for cell in guide[7]:
        cell.font = Font(color='FFFFFF', bold=True); cell.fill = PatternFill('solid', fgColor='087E8B')
    guide.sheet_view.showGridLines = False
    output = io.BytesIO(); book.save(output); book.close()
    return output.getvalue()


def register_import_routes(app, db, clock, supervisor):
    @app.get('/api/imports/template')
    def download_template(kind: Kind, user=Depends(supervisor)):
        with db.connect() as conn:
            regions = [dict(r) for r in conn.execute('SELECT * FROM regions WHERE active=1 ORDER BY id')]
        return Response(template(kind, regions), media_type=MIME,
                        headers={'Content-Disposition': f'attachment; filename="template-operasional-{kind}.xlsx"'})

    @app.post('/api/imports/upload', status_code=201)
    async def upload_workbook(request: Request, kind: Kind, filename: str = Query(min_length=1, max_length=120), user=Depends(supervisor)):
        content = bytearray()
        async for chunk in request.stream():
            content.extend(chunk)
            if len(content) > MAX_BYTES:
                raise HTTPException(413, 'Berkas maksimal 5 MB.')
        return await run_in_threadpool(upload, db, bytes(content), kind, filename, user, clock())

    @app.get('/api/imports')
    def history(user=Depends(supervisor)):
        with db.connect() as conn:
            rows = [dict(r) for r in conn.execute('SELECT id,kind,filename,state,revision,created_at,expires_at,result FROM import_batches WHERE actor_id=? ORDER BY created_at DESC,upload_order DESC LIMIT 20', (user['id'],))]
        for row in rows:
            row['total'] = json.loads(row.pop('result') or '{}').get('total', 0)
            if row['state'] not in {'committed', 'discarded'} and row['expires_at'] <= timestamp(clock()):
                row['state'] = 'expired'
        return rows

    @app.get('/api/imports/{batch_id}')
    def read_batch(batch_id: str, offset: int = Query(0, ge=0), limit: int = Query(25, ge=1, le=200), user=Depends(supervisor)):
        with db.connect() as conn:
            return batch_result(require_batch(conn, batch_id, user, clock(), allow_expired=True), clock(), offset, limit)

    @app.get('/api/imports/{batch_id}/layout')
    def read_layout(batch_id: str, sheet: str = Query(max_length=31), header_row: int = Query(1, ge=1, le=20), user=Depends(supervisor)):
        with db.connect() as conn:
            batch = require_batch(conn, batch_id, user, clock())
        if batch['state'] == 'committed':
            raise HTTPException(409, 'Batch sudah diimpor.')
        with workbook(batch['content']) as book:
            return sheet_layout(book, batch['kind'], sheet, header_row)

    @app.post('/api/imports/{batch_id}/preview')
    def preview_batch(batch_id: str, payload: PreviewInput, user=Depends(supervisor)):
        return preview(db, batch_id, payload, user, clock())

    @app.post('/api/imports/{batch_id}/commit')
    def commit_batch(batch_id: str, payload: CommitInput, user=Depends(supervisor)):
        return commit(db, batch_id, payload, user, clock())

    @app.post('/api/imports/{batch_id}/discard')
    def discard_batch(batch_id: str, user=Depends(supervisor)):
        now = clock()
        with db.connect(write=True) as conn:
            batch = require_batch(conn, batch_id, user, now, allow_expired=True)
            if batch['state'] == 'committed':
                raise HTTPException(409, 'Kasus yang sudah diimpor tidak dapat dibuang lewat draft.')
            if batch['state'] != 'discarded':
                conn.execute("UPDATE import_batches SET state='discarded',revision=revision+1,content=NULL,preview=NULL WHERE id=?", (batch_id,))
                audit(conn, user, 'import_discard', None, dict(batch_id=batch_id), now)
            return batch_result(require_batch(conn, batch_id, user, now, allow_expired=True), now)
