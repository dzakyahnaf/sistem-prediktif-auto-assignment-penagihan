import io
import json
import shutil
import tempfile
import unittest
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from pathlib import Path
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.imports import MAX_BYTES
from app.main import create_app
from app.security import password_hash
from app.seed import DEMO_PASSWORD


def xlsx(rows=None, headers=None, title='Data', prefix=False):
    book = Workbook(); sheet = book.active; sheet.title = title
    if prefix:
        sheet.append(['Judul workbook operasional'])
    sheet.append(headers or ['source_ref', 'customer', 'region_id', 'difficulty', 'amount', 'dpd'])
    for row in rows if rows is not None else [['S5-001', 'Alias Impor 001', 'SMG', 'Low', 1000000, 12]]:
        sheet.append(row)
    stream = io.BytesIO(); book.save(stream); book.close()
    return stream.getvalue()


class ImportTests(unittest.TestCase):
    def test_policy_change_invalidates_preview_then_commit_uses_current_rules(self):
        batch=self.preview(self.upload())
        p=self.client.get('/api/decisions',headers=self.headers).json()['policy']
        config={**p['config'],'collection_base':90,'medium_threshold':70}
        r=self.client.post('/api/decisions/policy',headers=self.headers,json=dict(config=config,expected_revision=p['revision'],reason='Validasi preview impor lintas kebijakan',request_key=uuid4().hex))
        self.assertEqual(r.status_code,200,r.text);self.assertEqual(self.commit(batch).status_code,409)
        fresh=self.preview(batch);result=self.commit(fresh);self.assertEqual(result.status_code,200,result.text)
        with self.app.state.db.connect() as c:
            row=c.execute("SELECT score_version,score FROM cases WHERE source='excel-operational'").fetchone()
        self.assertEqual(row['score_version'],'rules-policy-v1');self.assertEqual(row['score'],84.8)
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(); cls.base = str(Path(cls.temp.name) / 'base.db')
        with TestClient(create_app(cls.base)):
            pass

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def setUp(self):
        self.now = datetime(2026, 10, 28, 2, tzinfo=timezone.utc)
        self.path = str(Path(self.temp.name) / f'{uuid4().hex}.db'); shutil.copy2(self.base, self.path)
        self.app = create_app(self.path, lambda: self.now)
        self.client = TestClient(self.app).__enter__(); self.headers = self.login('supervisor')

    def tearDown(self):
        self.client.__exit__(None, None, None)

    def login(self, name):
        response = self.client.post('/api/auth/login', json={'username': name, 'password': DEMO_PASSWORD})
        response.raise_for_status()
        return {'Authorization': 'Bearer ' + response.json()['access_token']}

    def upload(self, content=None, kind='collection'):
        r = self.client.post('/api/imports/upload', params=dict(kind=kind, filename='fixture.xlsx'),
                             content=content if content is not None else xlsx(), headers=self.headers)
        r.raise_for_status(); return r.json()

    def preview(self, batch, mapping=None, **extra):
        meta = batch['metadata']
        layout = self.client.get(f'/api/imports/{batch["id"]}/layout', params=dict(sheet=meta['suggested_sheet'], header_row=meta['suggested_header_row']), headers=self.headers)
        layout.raise_for_status(); data = layout.json()
        response = self.client.post(f'/api/imports/{batch["id"]}/preview', headers=self.headers,
                                    json=dict(sheet=data['sheet'], header_row=data['header_row'], mapping=mapping if mapping is not None else data['mapping'], **extra))
        response.raise_for_status(); return response.json()

    def commit(self, batch, **extra):
        return self.client.post(f'/api/imports/{batch["id"]}/commit', json=dict(expected_revision=batch['revision'], **extra), headers=self.headers)

    def count(self, table='cases'):
        with self.app.state.db.connect() as db:
            return db.execute(f'SELECT COUNT(*) FROM {table}').fetchone()[0]

    def test_templates_empty_data_typed_examples_and_both_workflows(self):
        for kind in ['collection', 'origination']:
            response = self.client.get('/api/imports/template', params={'kind': kind}, headers=self.headers)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers['cache-control'], 'no-store')
            book = load_workbook(io.BytesIO(response.content))
            self.assertEqual(book.sheetnames, ['Data', 'Petunjuk'])
            self.assertEqual(book['Data'].max_row, 1)
            self.assertEqual(book['Data'].freeze_panes, 'A2')
            self.assertEqual(book['Data'].column_dimensions['A'].number_format, '@')
            book.close()
            empty = self.preview(self.upload(response.content, kind))
            self.assertEqual(empty['preview']['total'], 0)
            self.assertFalse(empty['preview']['can_commit'])
        original = self.count()
        collection = self.preview(self.upload())
        self.assertEqual(collection['preview']['items'][0]['data']['score'], 78.8)
        self.assertEqual(self.count(), original)
        result = self.commit(collection).json()
        cid = result['result']['items'][0]['case_id']
        self.assertEqual(result['result']['total'], 1)
        self.assertEqual(self.client.post('/api/assignments', json={'case_ids': [cid]}, headers=self.headers).json()['results'][0]['assigned'], True)
        headers = ['source_ref', 'customer', 'region_id', 'difficulty', 'amount', 'income', 'interest', 'documents_complete']
        origin = self.preview(self.upload(xlsx([['S5-001', 'Alias Pengajuan', 'Semarang', 'medium', 2000000, 5000000, 4, 'Ya']], headers), 'origination'))
        self.assertTrue(origin['preview']['items'][0]['data']['documents_complete'])
        self.assertTrue(self.commit(origin).is_success)  # Same source ID allowed in a different domain.
        self.assertEqual(self.count('import_case_refs'), 2)

    def test_invalid_rows_formulas_duplicates_and_numerical_types_block_everything(self):
        rows = [
            ['DUP', 'Alias Valid', 'SMG', 'Low', 1000, 0],
            ['dup', 'Alias Kedua', 'SMG', 'Low', 1000, 0],
            [123, 'Alias Angka', 'XXX', 'Unexpected', '1.000.000', -1],
            ['FORMULA', '=1+1', 'SMG', 'High', 1000.5, 0],
            ['BOOL', 'Alias Boolean', 'SMG', 'Low', True, 0],
        ]
        batch = self.preview(self.upload(xlsx(rows)))
        self.assertEqual(batch['preview']['invalid'], 5)
        errors = ' '.join(e for row in batch['preview']['items'] for e in row['errors'])
        for text in ['berulang', 'formula', 'Wilayah', 'bilangan bulat', 'teks']:
            self.assertIn(text, errors)
        before = self.count(); self.assertEqual(self.commit(batch).status_code, 409)
        self.assertEqual(self.count(), before)

    def test_mapping_header_sheet_ignored_columns_acknowledged(self):
        headers = ['Ref', 'Alias', 'Area', 'Kesulitan', 'Tagihan', 'DPD', 'Hasil Call']
        batch = self.upload(xlsx([['MAN-1', 'Alias Manual', 'sby', 'High', '2000', '25', 'Hasil historis tidak diimpor']], headers, title='Operasional', prefix=True))
        self.assertEqual(batch['metadata']['suggested_header_row'], 2)
        payload = dict(sheet='Operasional', header_row=2, mapping=dict(source_ref=1, customer=2, region_id=3, difficulty=4, amount=5, dpd=6))
        r = self.client.post(f'/api/imports/{batch["id"]}/preview', json=payload, headers=self.headers)
        r.raise_for_status(); batch = r.json()
        self.assertEqual(batch['preview']['ignored_columns'][0]['header'], 'Hasil Call')
        self.assertEqual(self.commit(batch).status_code, 422)
        self.assertEqual(self.commit(batch, acknowledge_ignored=True).status_code, 200)
        with self.app.state.db.connect() as db:
            case = dict(db.execute("SELECT * FROM cases WHERE source='excel-operational'").fetchone())
        self.assertEqual((case['region_id'], case['paid'], case['status'], case['cycle']), ('SBY', 0, 'new', 0))

    def test_atomic_conflict_after_preview_rolls_back_earlier_rows(self):
        first = self.preview(self.upload(xlsx([['SAME', 'Alias Pertama', 'SMG', 'Low', 1000, 0]])))
        second = self.preview(self.upload(xlsx([['UNIQUE', 'Alias Unik', 'SMG', 'Low', 1000, 0], ['SAME', 'Alias Konflik', 'SMG', 'Low', 1000, 0]])))
        self.commit(first).raise_for_status()
        before, audits = self.count(), self.count('audit')
        self.assertEqual(self.commit(second).status_code, 409)
        self.assertEqual((self.count(), self.count('audit')), (before, audits))
        again = self.preview(second)
        self.assertEqual((again['preview']['valid'], again['preview']['invalid']), (1, 1))
        with self.app.state.db.connect() as db:
            self.assertIsNone(db.execute("SELECT 1 FROM import_case_refs WHERE source_key='UNIQUE'").fetchone())

    def test_concurrent_commit_retries_create_once_and_audit_once(self):
        batch = self.preview(self.upload())
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda _: self.commit(batch), range(2)))
        self.assertEqual([r.status_code for r in results], [200, 200])
        self.assertEqual({r.json()['replayed'] for r in results}, {True, False})
        self.assertEqual(self.count('import_case_refs'), 1)
        with self.app.state.db.connect() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM audit WHERE action='import_commit'").fetchone()[0], 1)
        self.assertEqual(results[0].json()['result']['items'], results[1].json()['result']['items'])

    def test_revision_prevents_commit_of_old_preview_and_owner_role_scope(self):
        batch = self.preview(self.upload())
        wrong = dict(source_ref=1, customer=2, region_id=3, difficulty=4, amount=6, dpd=5)
        changed = self.preview(batch, wrong)
        self.assertFalse(changed['preview']['can_commit'])
        self.assertEqual(self.commit(batch).status_code, 409)
        with self.app.state.db.connect(write=True) as db:
            db.execute('INSERT INTO users VALUES (?,?,?,?,?)', ('U-SUP-2', 'supervisor2', password_hash(DEMO_PASSWORD), 'supervisor', None))
        second = self.login('supervisor2'); agent = self.login('smg-col-a')
        for method, path, data in [('GET', '/api/imports', None), ('GET', '/api/imports/template?kind=collection', None),
                ('GET', f'/api/imports/{batch["id"]}', None), ('POST', f'/api/imports/{batch["id"]}/commit', {'expected_revision': 1}),
                ('POST', f'/api/imports/{batch["id"]}/preview', {'sheet': 'Data', 'header_row': 1, 'mapping': wrong}),
                ('POST', '/api/imports/upload?kind=collection&filename=a.xlsx', None), ('POST', f'/api/imports/{batch["id"]}/discard', None)]:
            self.assertEqual(self.client.request(method, path, json=data, headers=agent).status_code, 403)
        self.assertEqual(self.client.get(f'/api/imports/{batch["id"]}', headers=second).status_code, 404)
        self.assertEqual(self.client.get('/api/imports', headers=second).json(), [])
        self.now += timedelta(hours=9)
        self.assertEqual(self.client.get('/api/imports', headers=self.headers).status_code, 401)

    def test_persistence_expiry_discard_and_idempotent_schema_upgrade(self):
        batch = self.preview(self.upload())
        self.client.__exit__(None, None, None)
        self.app = create_app(self.path, lambda: self.now)
        self.client = TestClient(self.app).__enter__()
        restored = self.client.get(f'/api/imports/{batch["id"]}', headers=self.headers).json()
        self.assertEqual(restored['preview'], batch['preview'])
        self.assertEqual(self.commit(restored).status_code, 200)
        discarded = self.upload(xlsx([['OTHER', 'Alias Lain', 'SMG', 'Low', 1000, 0]]))
        self.assertEqual(self.client.post(f'/api/imports/{discarded["id"]}/discard', headers=self.headers).json()['state'], 'discarded')
        self.assertEqual(self.client.get(f'/api/imports/{discarded["id"]}/layout?sheet=Data', headers=self.headers).status_code, 410)
        waiting = self.upload()
        self.now += timedelta(hours=25); self.headers = self.login('supervisor')
        self.assertEqual(self.client.get(f'/api/imports/{waiting["id"]}', headers=self.headers).json()['state'], 'expired')
        self.assertEqual(self.client.get(f'/api/imports/{waiting["id"]}/layout?sheet=Data', headers=self.headers).status_code, 410)
        self.assertTrue(self.commit(batch).json()['replayed'])  # A committed batch never expires.
        self.app.state.db.initialize()
        self.assertEqual(self.count('import_case_refs'), 1)

    def test_workbook_limits_malformed_macros_and_mapping_errors(self):
        for content, filename, expected in [(b'not zip', 'a.xlsx', 422), (b'not zip', 'a.xls', 422), (b'x' * (MAX_BYTES + 1), 'a.xlsx', 413)]:
            r = self.client.post('/api/imports/upload', params={'kind': 'collection', 'filename': filename}, content=content, headers=self.headers)
            self.assertEqual(r.status_code, expected)
        content = io.BytesIO(xlsx())
        with zipfile.ZipFile(content, 'a') as archive:
            archive.writestr('xl/vbaProject.bin', b'macro')
        self.assertEqual(self.client.post('/api/imports/upload?kind=collection&filename=a.xlsx', content=content.getvalue(), headers=self.headers).status_code, 422)
        batch = self.upload()
        for mapping in [dict(source_ref=1), dict(source_ref=True, customer=2, region_id=3, difficulty=4, amount=5),
                        dict(source_ref=1, customer=1, region_id=3, difficulty=4, amount=5),
                        dict(source_ref=1, customer=2, region_id=3, difficulty=4, amount=64)]:
            r = self.client.post(f'/api/imports/{batch["id"]}/preview', json=dict(sheet='Data', header_row=1, mapping=mapping), headers=self.headers)
            self.assertEqual(r.status_code, 422)
        self.assertEqual(self.count('import_case_refs'), 0)

    def test_two_thousand_rows_preview_pagination_commit_no_truncation(self):
        rows = [[f'BULK-{i:04}', f'Alias Bulk {i:04}', 'BDG', 'Low', 1000, 0] for i in range(2000)]
        batch = self.preview(self.upload(xlsx(rows)))
        self.assertEqual((batch['preview']['total'], batch['preview']['valid']), (2000, 2000))
        self.assertEqual(len(batch['preview']['items']), 25)
        last = self.client.get(f'/api/imports/{batch["id"]}?offset=1975&limit=25', headers=self.headers).json()
        self.assertEqual(last['preview']['items'][-1]['source_ref'], 'BULK-1999')
        self.assertEqual(self.commit(batch).json()['result']['total'], 2000)
        self.assertEqual(self.count('import_case_refs'), 2000)
        too_many = xlsx(rows + [['OVER', 'Alias Berlebih', 'BDG', 'Low', 1000, 0]])
        extra = self.upload(too_many)
        r = self.client.post(f'/api/imports/{extra["id"]}/preview', json=dict(sheet='Data', header_row=1, mapping=batch['preview']['mapping']), headers=self.headers)
        self.assertEqual(r.status_code, 413)
