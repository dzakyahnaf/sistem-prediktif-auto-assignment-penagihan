"""Prepare a new isolated database and XLSX fixtures for Sprint 5 browser UAT."""
import argparse
from contextlib import closing
import io
import json
from pathlib import Path
import runpy
import sqlite3
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'backend'))
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from app.db import Database
from app.imports import template

fingerprint = runpy.run_path(str(ROOT / 'scripts/prepare-pagination-uat.py'))['fingerprint']


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--database', type=Path, default=ROOT / 'runtime/imports-uat.db')
    parser.add_argument('--fixtures', type=Path, default=ROOT / 'runtime/import-fixtures')
    args = parser.parse_args()
    source = (ROOT / 'runtime/assignment.db').resolve(); target = args.database.resolve(); folder = args.fixtures.resolve()
    if source == target or target.exists() or folder.exists():
        parser.error('Choose new database and fixture paths; existing results are never overwritten.')
    original = fingerprint(source)
    target.parent.mkdir(parents=True, exist_ok=True); folder.mkdir(parents=True)
    with closing(sqlite3.connect(source.as_uri() + '?mode=ro', uri=True)) as live:
        with closing(sqlite3.connect(target)) as copy:
            live.backup(copy)
    Database(str(target)).initialize()
    with closing(sqlite3.connect(target)) as db:
        regions = [dict(id=r[0], name=r[1]) for r in db.execute('SELECT * FROM regions ORDER BY id')]
    report = dict(database=str(target), source_fingerprint=original, files={})

    def fixture(name, kind, rows, *, ignored=False, manual=False):
        book = load_workbook(io.BytesIO(template(kind, regions)))
        ws = book['Data']
        if manual:
            ws['E1'] = 'Saldo_Custom'
            ws.insert_rows(1); ws['A1'] = 'UAT S5 workbook operasional dengan judul'
        if ignored:
            col = ws.max_column + 1; ws.cell(1, col, 'Hasil_Call_Siklus_Berjalan')
        for row in rows:
            ws.append(row)
        if manual:
            ws.freeze_panes = 'A3'
        file = folder / name
        book.save(file); book.close()
        report['files'][name] = dict(path=str(file), kind=kind, rows=len(rows), manual_mapping=manual, ignored=ignored)

    fixture('penagihan-valid.xlsx', 'collection', [
        ['COL-S5-0001', 'UAT S5 Penagihan 001', 'SMG', 'Low', 100000, 0, 'Tidak dipakai untuk skor awal'],
        ['COL-S5-0002', 'UAT S5 Penagihan 002', 'SMG', 'Medium', 200000, 15, 'Tidak dipakai untuk skor awal'],
        ['COL-S5-0003', 'UAT S5 Penagihan 003', 'BDG', 'High', 300000, 80, 'Tidak dipakai untuk skor awal'],
    ], ignored=True)
    fixture('pengajuan-valid.xlsx', 'origination', [
        ['ORI-S5-0001', 'UAT S5 Pengajuan 001', 'SMG', 'Low', 1000000, 5000000, 4, 'Ya'],
        ['ORI-S5-0002', 'UAT S5 Pengajuan 002', 'SBY', 'High', 2000000, 6000000, 3, 'Tidak'],
    ])
    fixture('pemetaan-header-2.xlsx', 'collection', [
        ['MAP-S5-0001', 'UAT S5 Pemetaan 001', 'Semarang', 'Low', 400000, 3],
        ['MAP-S5-0002', 'UAT S5 Pemetaan 002', 'Bandung', 'High', 500000, 40],
    ], manual=True)
    fixture('preview-31-baris-invalid.xlsx', 'collection', [
        [f'CHECK-S5-{i:04}', f'UAT S5 Preview {i:04}', 'SMG' if i < 30 else 'UNKNOWN', 'Low', 1000, 0] for i in range(31)
    ])
    fixture('formula-duplikat-angka.xlsx', 'collection', [
        ['DUP-S5', 'UAT S5 Duplikat 1', 'SMG', 'Low', 1000, 0],
        ['dup-s5', 'UAT S5 Duplikat 2', 'SMG', 'Low', 1000, 0],
        [123, 'UAT S5 ID Angka', 'SMG', 'Low', '1.000.000', 0],
        ['FORM-S5', '=SUM(1,2)', 'SMG', 'Low', 1000, 0],
    ])
    for kind in ['collection', 'origination']:
        file = folder / f'template-kosong-{kind}.xlsx'; file.write_bytes(template(kind, regions))
    assert fingerprint(source) == original
    target.with_suffix('.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    print(json.dumps(report, indent=2))


if __name__ == '__main__':
    main()
